from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path("docs/brief/Рабочий_бриф_КосмоХакатон_2026.docx")

INK = "17212B"
MUTED = "66727D"
ACCENT = "0F4C5C"
PALE = "EAF3F5"
BORDER = "D9D9D9"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        node.append(el)
    tbl_pr.append(node)


def cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color=INK, size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)


def add_table(doc, headers, rows, widths, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        hdr.cells[i].width = Inches(widths[i])
        shade(hdr.cells[i], ACCENT)
        set_cell_text(hdr.cells[i], text, bold=True, color=WHITE, size=9)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = Inches(widths[i])
            if ridx % 2:
                shade(cells[i], PALE)
            alignment = alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], str(text), align=alignment)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([r_pr, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_link_item(doc, label, url, note):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    add_hyperlink(p, label, url)
    p.add_run(f" — {note}")


def add_bullet(doc, text, level=0, checked=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    if checked is not None:
        p.add_run("☐ " if not checked else "☒ ")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.add_run(text)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    old = p_pr.find(qn("w:pBdr"))
    if old is not None:
        p_pr.remove(old)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        p_bdr.append(el)
    p_pr.append(p_bdr)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.13

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_after = Pt(10)

    for name, size, before, after in (("Heading 1", 16, 16, 7), ("Heading 2", 12.5, 12, 5), ("Heading 3", 11, 9, 3)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(INK)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.text = "КОСМОХАКАТОН 2026   РАБОЧИЙ БРИФ"
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.add_run("Внутренний рабочий документ   •   ")
    footer.runs[0].font.name = "Aptos"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    page_number(footer)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run("РАБОЧАЯ ВЕРСИЯ   19 СЕНТЯБРЯ 2026")
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(ACCENT)

    title_p = doc.add_paragraph("Рабочий бриф команды КосмоХакатон 2026", style="Title")
    remove_paragraph_borders(title_p)
    subtitle = doc.add_paragraph("Прогнозирование внешних рисков и планирование ВКД на МКС")
    subtitle.paragraph_format.space_after = Pt(18)
    rr = subtitle.runs[0]
    rr.font.name = "Aptos Display"
    rr.font.size = Pt(14)
    rr.font.color.rgb = RGBColor.from_string(MUTED)

    doc.add_heading("Решение и текущий вывод", level=1)
    doc.add_paragraph(
        "Команда разрабатывает исследовательский сервис ВКД Риск для аналитика наземной группы. "
        "Сервис должен сопоставлять космическую погоду, орбиту МКС и второй механизм воздействия, "
        "сравнивать окна одинаковой длительности и объяснять рекомендацию через исходные данные."
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    lead = p.add_run("Главный вывод. ")
    lead.bold = True
    p.add_run(
        "Второй Colab уже задаёт сильный каркас демонстрации: поиск проекта, тесты, расчёт по двум линиям риска, "
        "ранжирование окон, отчёт и сборка сайта. Для конкурентной сдачи нужно закрыть четыре доказуемых пробела: "
        "исторический replay, живой пересчёт в интерфейсе, экспериментальное сравнение и полный пакет воспроизводимости."
    )

    add_table(
        doc,
        ["Область", "Что подтверждено", "Оценка состояния"],
        [
            ["Замысел", "Поддержка решения о времени ВКД на МКС", "Подтверждено постановкой"],
            ["Прототип", "Colab описывает расчёт, сайт и необязательный API", "Частично подтверждено"],
            ["Данные", "В коде названы NOAA SWPC и CelesTrak", "Живой запуск не проверен"],
            ["История", "Есть механизм Store as_of", "Загрузчик архивных выпусков отсутствует"],
            ["Интерфейс", "Статическая страница собирается", "Форма пока не пересчитывает"],
            ["Сдача", "Критерии и постановка доступны", "Ссылки на деплой и репозиторий не даны"],
        ],
        [1.25, 3.45, 1.75],
    )

    doc.add_heading("Основание документа", level=2)
    doc.add_paragraph(
        "Бриф собран из переписки от 18 сентября 2026 года, двух переданных Colab-блокнотов, постановки задачи "
        "и критериев оценки из рабочей папки. Содержимое Яндекс Доски не загрузилось в доступном браузере; "
        "поэтому доска учтена как канал управления, но её карточки не пересказаны."
    )

    page_break(doc)
    doc.add_heading("Хронология обмена", level=1)
    add_table(
        doc,
        ["Время", "Участник", "Сообщение или ресурс", "Рабочее значение"],
        [
            ["20:32:07", "ёжик в тумане", "Ссылка на Яндекс Телемост", "Канал синхронной встречи"],
            ["20:32:38", "agent 008", "Ищет место и планирует подключиться через пять минут", "Участник подключается с задержкой"],
            ["21:05:42", "анан", "Первый Google Colab", "Минимальный блокнот с подключением Google Диска"],
            ["21:05:42", "анан", "Яндекс Доска для менеджмента", "Предполагаемый источник задач и статусов"],
            ["21:05:42", "анан", "Второй Google Colab", "Болванка от Claude для запуска ВКД Риск"],
        ],
        [0.75, 1.0, 2.65, 2.05],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("Рабочие ссылки", level=1)
    add_link_item(doc, "Яндекс Телемост", "https://telemost.360.yandex.ru/j/2409243815", "встреча команды")
    add_link_item(doc, "Первый Colab", "https://colab.research.google.com/drive/1SEXwi7YT4WDif0hTUXbvz-TTg7F1qEVX?usp=sharing", "подключает Google Диск; содержательной логики почти нет")
    add_link_item(doc, "Яндекс Доска", "https://boards.yandex.ru/whiteboard/?hash=74b01f9008b12701a266597ddc9e08a6", "доска управления Космохак")
    add_link_item(doc, "Colab ВКД Риск", "https://colab.research.google.com/drive/1SW8n6bF4CPnSxljVq84Bo3PbjwzvLK2J?usp=sharing", "основной сценарий запуска и демонстрации")

    doc.add_heading("Ограничения источников", level=2)
    add_bullet(doc, "Первый Colab содержит только подключение Google Диска, заголовок New section и пустые ячейки.")
    add_bullet(doc, "Яндекс Доска открывается с названием Космохак, но полотно и карточки не отрисовались.")
    add_bullet(doc, "Код пакета eva risk в рабочей папке не найден; выводы о реализации основаны на сценарии второго Colab.")
    add_bullet(doc, "Факт успешного живого запуска, деплоя и доступности источников не подтверждён.")

    page_break(doc)
    doc.add_heading("Пользовательская задача", level=1)
    doc.add_paragraph(
        "Основной пользователь — аналитик наземной группы планирования ВКД. Он задаёт начало и длительность выхода, "
        "допустимый период переноса и режим анализа. Сервис должен помочь понять, какие внешние факторы относятся "
        "к выбранному окну, чем варианты отличаются и каких данных не хватает для уверенного выбора."
    )

    doc.add_heading("Обязательный сценарий", level=2)
    for text in [
        "Пользователь выбирает текущую обстановку или историческую дату в диапазоне с 1 мая по 30 июня 2024 года.",
        "Задаёт начало ВКД, длительность от 1 до 8 часов и период поиска вариантов до 24 часов.",
        "Сервис рассчитывает траекторию МКС по орбитальным элементам и показывает их эпоху и давность.",
        "Система оценивает космическую погоду и не менее одного другого механизма воздействия.",
        "Не менее двух окон одинаковой длительности сравниваются по воздействиям, времени пересечения и полноте данных.",
        "Для каждого предупреждения доступны значения, единицы, время, источник, правило, ограничения и уверенность.",
        "Пользователь меняет начало или длительность, получает пересчёт и сохраняет человекочитаемый и машиночитаемый результат.",
    ]:
        add_number(doc, text)

    doc.add_heading("Границы интерпретации", level=2)
    doc.add_paragraph(
        "Прототип оценивает внешнюю среду и условия на траектории. Он не определяет медицинскую допустимость выхода, "
        "индивидуальную дозу космонавта или вероятность повреждения скафандра без дополнительных моделей. "
        "Отсутствие данных нельзя показывать как отсутствие риска."
    )

    doc.add_heading("Что должно быть видно в результате", level=2)
    add_table(
        doc,
        ["Экран или блок", "Минимальное содержание"],
        [
            ["Запрос", "Режим, начало, длительность, период поиска и часовой пояс"],
            ["Траектория", "Содержательная сводка положения МКС, источник TLE или OMM, эпоха и давность"],
            ["Воздействия", "Временная картина факторов, текущие состояния и ожидаемые изменения"],
            ["Сравнение", "Несколько окон одной длительности, пересечения, полнота данных и компромиссы"],
            ["Доказательства", "Исходные значения, единицы, публикация, правило или модель и ограничения"],
            ["Экспорт", "Параметры, версии данных и алгоритма, рекомендация, отчёт и JSON или CSV"],
        ],
        [1.5, 4.95],
    )

    page_break(doc)
    doc.add_heading("Каркас прототипа из Colab", level=1)
    doc.add_paragraph(
        "Блокнот ВКД Риск ищет проект на Google Диске по пакету evarisk, устанавливает зависимости, запускает тесты, "
        "выполняет расчёт, собирает статический сайт и при необходимости поднимает FastAPI. Ниже указано то, "
        "что заявлено в блокноте; наличие и полнота самих модулей отдельно не проверены."
    )

    add_table(
        doc,
        ["Компонент", "Назначение по блокноту", "Что нужно подтвердить"],
        [
            ["sources", "Получение записей из реестра источников", "Живые ответы, единицы, частоты, ошибки и квоты"],
            ["provenance", "Хранилище записей, давность и отсечение as of", "Версии выпусков и строгая временная отсечка"],
            ["orbit", "Демонстрационный TLE и расчёт траектории", "Подходящие текущие и исторические элементы МКС"],
            ["risk spaceweather", "Оценка космической погоды", "Физический смысл порогов и прогноз не менее чем на 6 часов"],
            ["risk mmod", "Вторая линия риска", "Источник сближений или иная воспроизводимая модель"],
            ["risk fusion", "Объединение факторов", "Независимость сигналов и отсутствие двойного счёта"],
            ["windows", "Генерация, оценка и рекомендация окон", "Равнозначные варианты и недостаток оснований"],
            ["report", "Экспорт и текстовый отчёт", "Совпадение с интерфейсом и версии алгоритма"],
            ["site", "График воздействий и статическая страница", "Интерактивный пересчёт и состояния ошибок"],
            ["api", "POST plan и GET sources", "Валидация, изоляция запросов и стабильный деплой"],
        ],
        [1.35, 2.45, 2.65],
    )

    doc.add_heading("Параметры демонстрации", level=2)
    add_bullet(doc, "Начало по умолчанию — 11 мая 2024 года 00:00 UTC.")
    add_bullet(doc, "Длительность — 5,5 часа; поиск — 12 часов; шаг кандидатов — 120 минут.")
    add_bullet(doc, "Режимы — офлайн с синтетическим SEP событием и live с открытыми источниками.")
    add_bullet(doc, "Офлайн сценарий использует peak pfu 900 и демонстрационные шкалы S3, G4, R2.")
    add_bullet(doc, "В демонстрацию вручную добавлено сближение COSMOS 1408 DEB с TCA через 5,2 часа.")

    page_break(doc)
    doc.add_heading("Соответствие критериям", level=1)
    doc.add_paragraph(
        "Матрица ниже — предварительная оценка по доступным материалам. Статус Есть означает, что нужная механика "
        "явно описана в Colab. Он не заменяет проверку работающего приложения и исходного кода."
    )
    doc.add_heading("Отраслевые критерии", level=2)
    add_table(
        doc,
        ["Критерий", "Максимум", "Предварительный статус", "Главное подтверждение или пробел"],
        [
            ["О1 Факторы и задача", "8", "Частично", "Космическая погода и MMOD названы; обоснование связи с ВКД нужно оформить"],
            ["О2 Интерпретация", "8", "Частично", "Есть шкалы и оценки; нужны ограничения и аккуратный язык выводов"],
            ["О3 Сравнение окон", "8", "Есть", "candidate starts, score window и recommend заявлены"],
            ["О4 Доказательства", "8", "Частично", "Store и экспорт есть; переход от предупреждения к первичной записи надо показать"],
            ["О5 Удобство", "8", "Не готово", "Форма на странице не вызывает plan и не перерисовывает результат"],
            ["О6 Презентация", "5", "Не подтверждено", "Сценарий выступления и слайды не представлены"],
            ["О7 Дополнительные функции", "5", "Частично", "Парето фронт и несколько окон могут дать ценность при демонстрации"],
        ],
        [1.35, 0.65, 1.15, 3.3],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("Что уже можно показать", level=2)
    add_bullet(doc, "Сценарий запроса с изменяемыми началом, длительностью и периодом поиска.")
    add_bullet(doc, "Расчёт нескольких окон и рекомендацию через score window и recommend.")
    add_bullet(doc, "Временной график воздействий, текстовый отчёт и машиночитаемый payload.")

    page_break(doc)
    doc.add_heading("Технические критерии", level=1)
    add_table(
        doc,
        ["Критерий", "Максимум", "Предварительный статус", "Главное подтверждение или пробел"],
        [
            ["Т1 Данные", "7", "Частично", "Реестр источников заявлен; архивы, ошибки и происхождение надо подтвердить"],
            ["Т2 Орбита", "6", "Частично", "SGP4 и TLE заявлены; историческая орбита и эпоха требуют проверки"],
            ["Т3 Методы", "10", "Частично", "Два механизма и fusion заявлены; методику и двойной счёт надо объяснить"],
            ["Т4 Исторический режим", "8", "Пробел", "Нет загрузчика датированных выпусков NCEI и строгого replay"],
            ["Т5 Эксперименты", "6", "Пробел", "Нет сравнения с простым подходом, события и контрольного периода"],
            ["Т6 Надёжность", "5", "Частично", "Статусы и полнота есть; кеш, сбои и повторный расчёт не показаны"],
            ["Т7 Качество кода", "4", "Частично", "В блокноте заявлены 18 тестов; исходный код отсутствует в папке"],
            ["Т8 Документация", "4", "Пробел", "Нет подтверждённого репозитория, инструкции и сохранённых расчётов"],
        ],
        [1.35, 0.65, 1.15, 3.3],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("Самые дорогие пробелы", level=2)
    add_bullet(doc, "Т4 и Т5 вместе дают до 14 баллов и сейчас не подтверждены материалами.")
    add_bullet(doc, "О5 зависит от живого пользовательского пути, а не от статической страницы.")
    add_bullet(doc, "Т1, Т2 и Т3 требуют прослеживаемых данных и объяснения физического смысла, а не только работающего кода.")

    page_break(doc)
    doc.add_heading("План работ по приоритету", level=1)
    doc.add_paragraph(
        "Приоритет построен по влиянию на обязательные требования и баллы. Сначала команда должна получить "
        "воспроизводимый вертикальный сценарий, затем закрыть временную честность и исследование, после чего полировать демонстрацию."
    )

    doc.add_heading("Приоритет ноль Рабочий вертикальный сценарий", level=2)
    add_bullet(doc, "Зафиксировать ссылку на репозиторий, ветку и однозначную команду запуска.", checked=False)
    add_bullet(doc, "Запустить 18 тестов и сохранить лог успешного прогона.", checked=False)
    add_bullet(doc, "Подтвердить live получение NOAA SWPC и CelesTrak с временем публикации, единицами и давностью.", checked=False)
    add_bullet(doc, "Подключить форму сайта к POST plan и перерисовывать сравнение, предупреждения и источники.", checked=False)
    add_bullet(doc, "Развернуть доступную жюри версию и проверить её в чистом браузере.", checked=False)

    doc.add_heading("Приоритет один Баллы с высоким риском", level=2)
    add_bullet(doc, "Реализовать исторический запрос на произвольную дату с 1 мая по 30 июня 2024 года.", checked=False)
    add_bullet(doc, "Для космической погоды загрузить датированные выпуски и соблюдать cutoff публикации.", checked=False)
    add_bullet(doc, "Получать исторические орбитальные элементы или честно маркировать геометрию как реконструкцию.", checked=False)
    add_bullet(doc, "Подключить источник сближений или документировать второй механизм с воспроизводимыми данными.", checked=False)
    add_bullet(doc, "Различать нет воздействия, нет данных, данные устарели и источник недоступен.", checked=False)

    doc.add_heading("Приоритет два Исследование и доказательства", level=2)
    add_bullet(doc, "Выбрать выраженное событие и контрольный период без события, обосновать даты источниками.", checked=False)
    add_bullet(doc, "Сравнить метод с простым базовым вариантом на одинаковых окнах.", checked=False)
    add_bullet(doc, "Оценить заблаговременность, пропуски, ложные предупреждения и устойчивость выбора окна.", checked=False)
    add_bullet(doc, "Сохранить входные записи, параметры, cutoff, версии данных, версию алгоритма и результат.", checked=False)

    doc.add_heading("Приоритет три Пакет сдачи", level=2)
    add_bullet(doc, "README с архитектурой, источниками, ограничениями, установкой и двумя примерами запуска.", checked=False)
    add_bullet(doc, "Примеры сохранённых расчётов в читаемом формате и JSON или CSV.", checked=False)
    add_bullet(doc, "Презентация с проблемой, двумя механизмами, методом, экспериментом, демо и ограничениями.", checked=False)
    add_bullet(doc, "Короткий резервный сценарий демонстрации без сети на сохранённых данных.", checked=False)

    page_break(doc)
    doc.add_heading("Управление командой", level=1)
    doc.add_paragraph(
        "Из переписки нельзя достоверно восстановить роли, владельцев задач и дедлайн. Ниже — структура для заполнения "
        "на доске. Она превращает технические пробелы в проверяемые результаты."
    )
    add_table(
        doc,
        ["Поток", "Результат", "Владелец", "Статус"],
        [
            ["Данные", "Live и архивные источники с provenance", "Назначить", "Не подтверждено"],
            ["Методика", "Пороговые правила, fusion и ограничения", "Назначить", "Частично"],
            ["Исторический режим", "Replay с cutoff и версиями", "Назначить", "Не готово"],
            ["Интерфейс", "Форма, сравнение окон и доказательства", "Назначить", "Частично"],
            ["Эксперименты", "Событие, контроль, baseline и метрики", "Назначить", "Не готово"],
            ["Инфраструктура", "Репозиторий, тесты, деплой и резервное демо", "Назначить", "Не подтверждено"],
            ["Подача", "README, примеры, презентация и речь", "Назначить", "Не готово"],
        ],
        [1.35, 3.15, 0.95, 1.0],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )

    doc.add_heading("Рекомендуемый ритм", level=2)
    add_bullet(doc, "На доске каждая карточка должна иметь владельца, проверяемый результат и ссылку на доказательство.")
    add_bullet(doc, "Два раза в день команда проходит только блокеры и изменения статуса, а не пересказывает всю работу.")
    add_bullet(doc, "После каждого значимого изменения прогоняются тесты и сохраняется один контрольный расчёт.")
    add_bullet(doc, "Демо репетируется на live сценарии и на локально сохранённых данных.")

    doc.add_heading("Вопросы на ближайший созвон", level=2)
    add_number(doc, "Где находится репозиторий eva risk и какая версия соответствует Colab")
    add_number(doc, "Кто владеет историческим режимом и какие архивы уже проверены")
    add_number(doc, "Какой источник выбран для MMOD или сближений и допустима ли его лицензия")
    add_number(doc, "Где будет развёрнуто решение и кто проверяет доступ для жюри")
    add_number(doc, "Какие даты выбраны для события и контрольного периода")
    add_number(doc, "Кто собирает презентацию и какой сценарий показа укладывается во время")

    page_break(doc)
    doc.add_heading("Чек лист готовности", level=1)
    doc.add_heading("Продукт", level=2)
    for text in [
        "Запрос на 1–8 часов и поиск до 24 часов работает",
        "Сравниваются минимум два окна одинаковой длительности",
        "Изменение начала и длительности запускает пересчёт",
        "Все времена показаны с часовым поясом и в основном формате UTC",
        "Предупреждение раскрывается до значений, источника, модели и ограничений",
        "Экспорт совпадает с интерфейсом и содержит версии данных и алгоритма",
    ]:
        add_bullet(doc, text, checked=False)

    doc.add_heading("Данные и временная честность", level=2)
    for text in [
        "Текущие данные обновляются автоматически и вручную",
        "Видны время публикации, время получения, эпоха и давность",
        "Сбой или пропуск не превращается в зелёный статус",
        "Исторический режим работает на новой дате из заданного диапазона",
        "Replay использует только сведения до cutoff",
        "Поздние сведения применяются только для проверки",
    ]:
        add_bullet(doc, text, checked=False)

    doc.add_heading("Исследование и демонстрация", level=2)
    for text in [
        "Два разных механизма воздействия обоснованы",
        "Есть сравнение с простым базовым подходом",
        "Показаны выраженное событие и контрольный период",
        "Ограничения и границы применимости сформулированы без заявлений о безопасности",
        "Решение развёрнуто и доступно жюри",
        "Репозиторий, инструкция, примеры расчётов и презентация готовы",
    ]:
        add_bullet(doc, text, checked=False)

    doc.add_heading("Источники", level=1)
    add_bullet(doc, "Постановка задачи Прогнозирование внешних рисков и планирование ВКД на МКС, 7 страниц.")
    add_bullet(doc, "Критерии оценки Прогнозирование внешних рисков и планирование ВКД на МКС, 3 страницы.")
    add_bullet(doc, "Переписка команды от 18 сентября 2026 года.")
    add_bullet(doc, "Google Colab 1SEXwi7YT4WDif0hTUXbvz T Tg7F1qEVX, минимальная заготовка.")
    add_bullet(doc, "Google Colab 1SW8n6bF4CPnSxljVq84Bo3PbjwzvLK2J, сценарий запуска ВКД Риск.")

    core = doc.core_properties
    core.title = "Рабочий бриф команды КосмоХакатон 2026"
    core.subject = "Прогнозирование внешних рисков и планирование ВКД на МКС"
    core.author = "Команда КосмоХакатон"
    core.keywords = "КосмоХакатон, ВКД, МКС, риски, планирование"

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
